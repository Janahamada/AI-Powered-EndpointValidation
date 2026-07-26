"""
Report generation: executive-summary and per-endpoint PDFs (ReportLab) and a
findings workbook (OpenPyXL). All functions return raw bytes so the API layer
can stream them without touching the filesystem.
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.orm import Session

from app.config import settings
from app.services import compliance_service as cs

# --- Palette -------------------------------------------------------------- #
_BRAND = colors.HexColor("#1e3a8a")
_STATUS_COLORS = {
    "PASS": colors.HexColor("#16a34a"),
    "WARNING": colors.HexColor("#d97706"),
    "FAIL": colors.HexColor("#dc2626"),
    "NO_DATA": colors.HexColor("#6b7280"),
}
_SEVERITY_FILL = {
    "critical": "FFDC2626",
    "high": "FFEA580C",
    "medium": "FFD97706",
    "low": "FF64748B",
}


def _styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "H1Brand", parent=styles["Heading1"], textColor=_BRAND, spaceAfter=4
        )
    )
    styles.add(
        ParagraphStyle(
            "Sub", parent=styles["Normal"], textColor=colors.HexColor("#475569"),
            fontSize=9, spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            "SectionH", parent=styles["Heading2"], textColor=_BRAND,
            fontSize=12, spaceBefore=10, spaceAfter=6, alignment=TA_LEFT,
        )
    )
    return styles


def _header(styles, title: str, subtitle: str) -> list:
    return [
        Paragraph(settings.APP_NAME, styles["Sub"]),
        Paragraph(title, styles["H1Brand"]),
        Paragraph(subtitle, styles["Sub"]),
        Spacer(1, 6),
    ]


# --------------------------------------------------------------------------- #
# Fleet executive summary PDF
# --------------------------------------------------------------------------- #
def fleet_report_pdf(db: Session) -> bytes:
    summary = cs.build_dashboard(db)
    styles = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
        leftMargin=16 * mm, rightMargin=16 * mm, title="Fleet Assurance Report",
    )
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    story = _header(
        styles, "Fleet Assurance Executive Summary",
        f"Generated {now} • {summary.total_endpoints} endpoints assessed",
    )

    # KPI table
    kpis = [
        ["Overall Assurance Score", f"{summary.overall_compliance_score}%"],
        ["Total Endpoints", str(summary.total_endpoints)],
        ["Critical Findings", str(summary.critical_findings)],
        ["High-Risk Findings", str(summary.high_risk_findings)],
    ]
    kpi_table = Table(kpis, colWidths=[80 * mm, 40 * mm])
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2ff")),
                ("TEXTCOLOR", (0, 0), (0, -1), _BRAND),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c7d2fe")),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story += [Paragraph("Key Metrics", styles["SectionH"]), kpi_table, Spacer(1, 8)]

    # Per-control compliance
    story.append(Paragraph("Per-Control Compliance", styles["SectionH"]))
    ctrl_rows = [["Control", "Compliance", "Pass", "Warn", "Fail", "No Data"]]
    for pc in summary.per_control:
        ctrl_rows.append(
            [pc.label, f"{pc.compliance_rate}%", str(pc.pass_count),
             str(pc.warning_count), str(pc.fail_count), str(pc.no_data_count)]
        )
    ctrl_table = Table(ctrl_rows, colWidths=[60 * mm, 26 * mm, 18 * mm, 18 * mm, 18 * mm, 22 * mm])
    ctrl_table.setStyle(_grid_style())
    story += [ctrl_table, Spacer(1, 8)]

    # Top findings
    story.append(Paragraph("Top Findings (most severe)", styles["SectionH"]))
    find_rows = [["Severity", "Control", "Field", "Description"]]
    for f in summary.top_findings:
        find_rows.append([f.severity.upper(), f.control_type, f.field, _clip(f.description, 60)])
    if len(find_rows) == 1:
        find_rows.append(["—", "—", "—", "No findings. Fleet is fully compliant."])
    find_table = Table(find_rows, colWidths=[22 * mm, 26 * mm, 40 * mm, 74 * mm])
    find_table.setStyle(_grid_style(severity_col=0))
    story.append(find_table)

    doc.build(story)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Per-endpoint PDF
# --------------------------------------------------------------------------- #
def endpoint_report_pdf(db: Session, hostname: str) -> bytes | None:
    from app.services import remediation_service  # local import avoids a cycle

    ev = cs.evaluate_one(db, hostname)
    if ev is None:
        return None
    detail = cs.to_detail(ev)
    recommendations = remediation_service.build_recommendations(ev.findings)
    recommendations_text = remediation_service.render_text(hostname, recommendations)
    styles = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
        leftMargin=16 * mm, rightMargin=16 * mm, title=f"Endpoint Report — {hostname}",
    )
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    story = _header(
        styles, f"Endpoint Assurance Report",
        f"{hostname} • Generated {now}",
    )

    a = detail.asset
    meta = [
        ["Hostname", a.hostname, "Status", detail.status],
        ["IP Address", a.ip_address, "Compliance", f"{detail.compliance_score}%"],
        ["Operating System", a.operating_system, "Owner", a.business_owner],
    ]
    meta_table = Table(meta, colWidths=[32 * mm, 58 * mm, 26 * mm, 46 * mm])
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, 0), (0, -1), _BRAND),
                ("TEXTCOLOR", (2, 0), (2, -1), _BRAND),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story += [meta_table, Spacer(1, 8)]

    # Per-control status
    story.append(Paragraph("Control Validation", styles["SectionH"]))
    ctrl_rows = [["Control", "Status", "Score", "Passed"]]
    for c in detail.controls:
        ctrl_rows.append(
            [c.label, c.status, f"{c.score}%", f"{c.passed_fields}/{c.total_fields}"]
        )
    ctrl_table = Table(ctrl_rows, colWidths=[70 * mm, 30 * mm, 30 * mm, 32 * mm])
    ctrl_table.setStyle(_grid_style(status_col=1))
    story += [ctrl_table, Spacer(1, 8)]

    # Findings
    story.append(Paragraph("Findings", styles["SectionH"]))
    if detail.findings:
        rows = [["Severity", "Control", "Field", "Expected", "Actual"]]
        for f in detail.findings:
            rows.append(
                [f.severity.upper(), f.control_type, f.field,
                 _clip(str(f.expected), 22), _clip(str(f.actual), 22)]
            )
        ft = Table(rows, colWidths=[22 * mm, 26 * mm, 46 * mm, 34 * mm, 34 * mm])
        ft.setStyle(_grid_style(severity_col=0))
        story.append(ft)
    else:
        story.append(Paragraph("No findings — this endpoint is fully compliant.", styles["Normal"]))

    # Recommendations (deterministic, grounded, cited)
    if recommendations:
        story.append(Paragraph("Grounded Recommendations", styles["SectionH"]))
        for line in recommendations_text.split("\n"):
            story.append(
                Paragraph(_escape(line) if line.strip() else "&nbsp;", styles["Normal"])
            )

    doc.build(story)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Findings workbook (Excel)
# --------------------------------------------------------------------------- #
def findings_workbook(db: Session) -> bytes:
    evaluated = cs.evaluate_all(db)
    wb = Workbook()

    # Sheet 1: all findings
    ws = wb.active
    ws.title = "Findings"
    headers = ["Hostname", "IP", "Owner", "OS", "Control", "Field",
               "Severity", "Expected", "Actual", "Description"]
    _write_header(ws, headers)
    row_idx = 2
    for ev in evaluated:
        summ = cs.to_summary(ev)
        for f in ev.findings:
            ws.cell(row=row_idx, column=1, value=summ.hostname)
            ws.cell(row=row_idx, column=2, value=summ.ip_address)
            ws.cell(row=row_idx, column=3, value=summ.business_owner)
            ws.cell(row=row_idx, column=4, value=summ.operating_system)
            ws.cell(row=row_idx, column=5, value=f.control_type)
            ws.cell(row=row_idx, column=6, value=f.field)
            sev = ws.cell(row=row_idx, column=7, value=f.severity.upper())
            fill = _SEVERITY_FILL.get(f.severity)
            if fill:
                sev.fill = PatternFill("solid", fgColor=fill)
                sev.font = Font(color="FFFFFFFF", bold=True)
            ws.cell(row=row_idx, column=8, value=str(f.expected))
            ws.cell(row=row_idx, column=9, value=str(f.actual))
            ws.cell(row=row_idx, column=10, value=f.description)
            row_idx += 1
    _autosize(ws, headers)
    ws.freeze_panes = "A2"

    # Sheet 2: endpoint summary
    ws2 = wb.create_sheet("Endpoints")
    headers2 = ["Hostname", "IP", "Owner", "OS", "Status", "Score",
                "Critical", "High", "Total Findings"]
    _write_header(ws2, headers2)
    for i, ev in enumerate(evaluated, start=2):
        s = cs.to_summary(ev)
        for col, val in enumerate(
            [s.hostname, s.ip_address, s.business_owner, s.operating_system,
             s.status, s.compliance_score, s.critical_findings, s.high_findings,
             s.total_findings],
            start=1,
        ):
            ws2.cell(row=i, column=col, value=val)
    _autosize(ws2, headers2)
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _grid_style(severity_col: int | None = None, status_col: int | None = None) -> TableStyle:
    style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), _BRAND),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("PADDING", (0, 0), (-1, -1), 4),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
    )
    return style


def _write_header(ws, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF1E3A8A")
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, headers: list[str]) -> None:
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        longest = max(
            [len(str(headers[col_idx - 1]))]
            + [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 60)


def _clip(text: str, n: int) -> str:
    return text if len(text) <= n else text[: n - 1] + "…"


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
